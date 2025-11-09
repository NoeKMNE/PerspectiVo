from PySide6.QtWidgets import *
from PySide6.QtCore import *
from PySide6.QtGui import *
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from openpyxl import Workbook




class ModernCard(QFrame):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setFrameStyle(QFrame.Box)
        self.setStyleSheet("""
            QFrame {
                background-color: white;
                border: 1px solid #E5E7EB;
                border-radius: 6px;
                padding: 6px;
                margin: 2px;
            }
        """)

class ModernButton(QPushButton):
    def __init__(self, text, primary=False, parent=None):
        super().__init__(text, parent)
        if primary:
            self.setStyleSheet("""
                QPushButton {
                    background-color: #4F46E5;
                    color: white;
                    border: none;
                    border-radius: 5px;
                    padding: 4px 8px;
                    font-weight: bold;
                    font-size: 11px;
                }
                QPushButton:hover {
                    background-color: #4338CA;
                }
                QPushButton:pressed {
                    background-color: #3730A3;
                }
            """)
        else:
            self.setStyleSheet("""
                QPushButton {
                    background-color: #F3F4F6;
                    color: #374151;
                    border: 1px solid #D1D5DB;
                    border-radius: 5px;
                    padding: 4px 8px;
                    font-size: 11px;
                }
                QPushButton:hover {
                    background-color: #E5E7EB;
                }
            """)

class SidebarButton(QPushButton):
    def __init__(self, text, icon, parent=None):
        super().__init__(parent)
        self.setText(f"{text}")
        self.setIcon(icon)
        self.setCheckable(True)
        self.setStyleSheet("""
            QPushButton {
                text-align: left;
                padding: 6px 8px;
                margin: 1px 0;
                border: none;
                border-radius: 5px;
                font-size: 11px;
                background-color: transparent;
                color: #6B7280;
            }
            QPushButton:checked {
                background-color: #EEF2FF;
                color: #4F46E5;
                font-weight: bold;
            }
            QPushButton:hover {
                background-color: #F9FAFB;
            }
        """)


class MembersPage(QWidget):
    members_changed = Signal()
    def __init__(self, membres_mgr):
        super().__init__()
        
        self.membres_mgr = membres_mgr
        self.init_ui()
        self.load_sample_data()

    def init_ui(self):
        layout = QVBoxLayout(self)
        layout.setSpacing(20)
        layout.setContentsMargins(24, 24, 24, 24)

        # En-tête
        header_layout = QHBoxLayout()

        title = QLabel("Membres")
        title.setStyleSheet("font-size: 24px; font-weight: bold; color: #1F2937;")
        header_layout.addWidget(title)

        header_layout.addStretch()

        add_member_btn = ModernButton("+ Ajouter un membre", primary=True)
        add_member_btn.clicked.connect(self.add_member)
        header_layout.addWidget(add_member_btn)

        export_btn = ModernButton("Exporter en PDF")
        export_btn.clicked.connect(self.export_to_pdf)
        header_layout.addWidget(export_btn)

        # Nouveau bouton Excel
        export_excel_btn = ModernButton("Exporter en Excel")
        export_excel_btn.clicked.connect(self.export_to_excel)
        header_layout.addWidget(export_excel_btn)

        layout.addLayout(header_layout)

        # Barre de recherche
        search_layout = QHBoxLayout()
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("Rechercher un membre...")
        self.search_input.setFixedHeight(38)
        self.search_input.setStyleSheet("""
            QLineEdit {
                padding: 2px;
                border: 1px solid #D1D5DB;
                border-radius: 8px;
                font-size: 14px;
                background-color: white;
            }
            QLineEdit:focus {
                border-color: #4F46E5;
                outline: none;
            }
        """)
        self.search_input.textChanged.connect(self.search_members)
        search_layout.addWidget(self.search_input)
        layout.addLayout(search_layout)

        # Tableau des membres
        self.members_table = QTableWidget()
        self.members_table.setColumnCount(7)
        self.members_table.setHorizontalHeaderLabels([
            "Nom", "Prénoms", "Contact", "Email", "École", "Filière", "Taux de présence"
        ])

        # Style du tableau
        self.members_table.setStyleSheet("""
            QTableWidget {
                background-color: white;
                border: 1px solid #E5E7EB;
                border-radius: 12px;
                gridline-color: #F3F4F6;
                selection-background-color: #EEF2FF;
            }
            QTableWidget::item {
                padding: 12px 8px;
                border-bottom: 1px solid #F3F4F6;
                color: #374151;
            }
            QTableWidget::item:selected {
                background-color: #EEF2FF;
                color: #4F46E5;
            }
            QHeaderView::section {
                background-color: #F9FAFB;
                padding: 12px 8px;
                border: none;
                border-bottom: 2px solid #E5E7EB;
                font-weight: bold;
                color: #374151;
                text-align: left;
            }
        """)

        # Configuration du tableau
        header = self.members_table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.Stretch)
        self.members_table.setAlternatingRowColors(True)
        self.members_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.members_table.setShowGrid(True)
        self.members_table.verticalHeader().setVisible(False)

        layout.addWidget(self.members_table)

    def search_members(self):
        search_text = self.search_input.text().lower()
        for row in range(self.members_table.rowCount()):
            show_row = False
            for col in range(self.members_table.columnCount()):
                item = self.members_table.item(row, col)
                if item and search_text in item.text().lower():
                    show_row = True
                    break
            self.members_table.setRowHidden(row, not show_row)

    def load_sample_data(self):
        membres = self.membres_mgr.obtenir_tous_membres()
        self.members_table.setRowCount(len(membres))
        for row, membre in enumerate(membres):
            data = [
                membre['nom'], membre['prenoms'], membre['contact'],
                membre['email'], membre['ecole'], membre['filiere'],
                f"{self.membres_mgr.calculer_taux_presence(membre['id'])}%"
            ]
            for col, value in enumerate(data):
                item = QTableWidgetItem(str(value))
                item.setFlags(item.flags() & ~Qt.ItemIsEditable)
                if col == 6:
                    percentage = float(value.replace('%', ''))
                    if percentage >= 90:
                        item.setBackground(QColor(220, 252, 231))
                        item.setForeground(QColor(22, 101, 52))
                    elif percentage >= 80:
                        item.setBackground(QColor(254, 249, 195))
                        item.setForeground(QColor(120, 113, 108))
                    else:
                        item.setBackground(QColor(254, 226, 226))
                        item.setForeground(QColor(153, 27, 27))

                self.members_table.setItem(row, col, item)

    def add_member(self):
        dialog = AddMemberDialog(self)
        if dialog.exec() == QDialog.Accepted:
            try:
                self.membres_mgr.ajouter_membre(
                    nom=dialog.inputs['nom_input'].text(),
                    prenoms=dialog.inputs['prenoms_input'].text(),
                    contact=dialog.inputs['contact_input'].text(),
                    email=dialog.inputs['email_input'].text(),
                    residence=dialog.inputs['residence_input'].text(),
                    ecole=dialog.inputs['ecole_input'].text(),
                    filiere=dialog.inputs['filiere_input'].text()
                )
                self.load_sample_data()  # Recharger la table
                QMessageBox.information(self, "Succès", "Membre ajouté avec succès!")
            except ValueError as e:
                QMessageBox.warning(self, "Erreur", str(e))
                
    def export_to_pdf(self):
        from PySide6.QtWidgets import QFileDialog, QMessageBox
    
        filename, _ = QFileDialog.getSaveFileName(self, "Exporter en PDF", "membres.pdf", "PDF Files (*.pdf)")
        if not filename:
            return
    
        membres = self.membres_mgr.obtenir_tous_membres()
        data = [
            ["Nom", "Prénoms", "Contact", "Email", "École", "Filière", "Taux de présence"]
        ]
        for membre in membres:
            data.append([
                membre['nom'],
                membre['prenoms'],
                membre['contact'],
                membre['email'],
                membre['ecole'],
                membre['filiere'],
                f"{self.membres_mgr.calculer_taux_presence(membre['id'])}%"
            ])
    
        doc = SimpleDocTemplate(filename, pagesize=A4)
        styles = getSampleStyleSheet()
        elements = []
    
        title = Paragraph("Liste des Membres", styles['Title'])
        elements.append(title)
        elements.append(Spacer(1, 12))
    
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor("#F3F4F6")),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
            ('FONTSIZE', (0, 1), (-1, -1), 10),
        ]))
        elements.append(table)
    
        doc.build(elements)
    
        QMessageBox.information(self, "Export PDF", "Exportation réussie !")

    def export_to_excel(self):
        from PySide6.QtWidgets import QFileDialog, QMessageBox

        filename, _ = QFileDialog.getSaveFileName(self, "Exporter en Excel", "membres.xlsx", "Excel Files (*.xlsx)")
        if not filename:
            return

        membres = self.membres_mgr.obtenir_tous_membres()
        wb = Workbook()
        ws = wb.active
        ws.title = "Membres"

        # En-tête
        headers = ["Nom", "Prénoms", "Contact", "Email", "École", "Filière", "Taux de présence"]
        ws.append(headers)

        # Données
        for membre in membres:
            ws.append([
                membre['nom'],
                membre['prenoms'],
                membre['contact'],
                membre['email'],
                membre['ecole'],
                membre['filiere'],
                f"{self.membres_mgr.calculer_taux_presence(membre['id'])}%"
            ])

        try:
            wb.save(filename)
            QMessageBox.information(self, "Export Excel", "Exportation réussie !")
        except Exception as e:
            QMessageBox.critical(self, "Erreur", f"Erreur lors de l'exportation : {e}")


class AddMemberDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Ajouter un nouveau membre")
        self.setFixedSize(450, 610)
        self.setModal(True)
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout(self)
        layout.setSpacing(16)
        layout.setContentsMargins(24, 24, 24, 24)

        # Titre
        title = QLabel("Nouveau membre")
        title.setStyleSheet("font-size: 16px; font-weight: bold; color: #1F2937; margin-bottom: 12px;")
        title.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(title)

        # Champs du formulaire
        fields = [
            ("Nom *", "nom_input"),
            ("Prénoms *", "prenoms_input"),
            ("Contact *", "contact_input"),
            ("Email ", "email_input"),
            ("Résidence", "residence_input"),
            ("École", "ecole_input"),
            ("Filière", "filiere_input")
        ]

        self.inputs = {}
        for label_text, input_name in fields:
            label = QLabel(label_text)
            label.setStyleSheet("font-weight: bold; color: #374151; margin-bottom: 4px;")

            input_field = QLineEdit()
            input_field.setFixedHeight(34)
            input_field.setStyleSheet("""
                QLineEdit {
                    padding: 10px 12px;
                    border: 1px solid #D1D5DB;
                    border-radius: 6px;
                    font-size: 14px;
                    background-color: white;
                }
                QLineEdit:focus {
                    border-color: #4F46E5;
                    outline: none;
                }
            """)

            # Ajouter des placeholders
            placeholders = {
                "nom_input": "Ex: KOUASSI",
                "prenoms_input": "Ex: Jean-Baptiste",
                "contact_input": "Ex: 0123456789",
                "email_input": "Ex: jean@email.com",
                "residence_input": "Ex: 11 05",
                "ecole_input": "Ex: INPHB",
                "filiere_input": "Ex: BCPST"
            }
            input_field.setPlaceholderText(placeholders.get(input_name, ""))

            self.inputs[input_name] = input_field
            layout.addWidget(label)
            layout.addWidget(input_field)

        layout.addStretch()

        # Note sur les champs obligatoires
        note = QLabel("* Champs obligatoires")
        note.setStyleSheet("color: #DC2626; font-size: 12px; font-style: italic;")
        layout.addWidget(note)

        # Boutons
        buttons_layout = QHBoxLayout()
        buttons_layout.setSpacing(12)

        cancel_btn = ModernButton("Annuler")
        cancel_btn.clicked.connect(self.reject)

        save_btn = ModernButton("Ajouter le membre", primary=True)
        save_btn.clicked.connect(self.validate_and_accept)

        buttons_layout.addWidget(cancel_btn)
        buttons_layout.addWidget(save_btn)
        layout.addLayout(buttons_layout)

    def validate_and_accept(self):
        required_fields = ['nom_input', 'prenoms_input', 'contact_input']
        field_names = {
            'nom_input': 'Nom',
            'prenoms_input': 'Prénoms',
            'contact_input': 'Contact'
        }

        for field in required_fields:
            if not self.inputs[field].text().strip():
                QMessageBox.warning(self, "Erreur", f"Le champ '{field_names[field]}' est obligatoire.")
                self.inputs[field].setFocus()
                return

        # Validation de l'email
        
        email = self.inputs['email_input'].text().strip()
        if email != "":
            if '@' not in email or '.' not in email:
                QMessageBox.warning(self, "Erreur", "Veuillez saisir une adresse email valide.")
                self.inputs['email_input'].setFocus()
                return

            self.accept()
